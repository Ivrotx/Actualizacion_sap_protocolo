import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
import os
import re
import unicodedata
import tempfile

# ==========================================
# COLUMNAS REQUERIDAS POR SAP
# ==========================================
COLUMNAS_PROTOCOLO = [
    'ResControlProd', 'GrupoPlanificador', 'Material', 'Descripcion', 
    'GrupoReceta', 'Version', 'Ancho', 'Perimetro', 'PorcEficiencia', 
    'NombreRecurso', 'EquipoXRecurso', 'Consecutivo', 'Velocidad', 
    'Alistamiento', 'OperarioXRecurso', '%Enlace', '%Recorte', 
    '%Sobrepeso', 'AfectaEnlace'
]

COLUMNAS_ESPECIFICACIONES = [
    'MATERIAL', 'Peso GTA cruda', 'Peso GTA Horneada', 'PESO GTA ROCEADA (gr)', 
    'Gtas/Sanduche', 'Base tiempo/min', 'Cant.crema sanduche', 'Galleta /paqueticos', 
    'PAQUETICOS UNIDAD', 'UNIDADES CORRUGADO', 'LARGO', 'ANCHO', 'CALIBRE', 
    'PESO N.IMPRESO UV(KG)', 'PESO PASTA CRUDA'
]

MAPEO_BASE = {
    'PorcEficiencia': ['porceficiencia', 'eficiencia', '% eficiencia'],
    'GrupoPlanificador': ['grupo planificador', 'grpplanif', 'planificador'],
    'NombreRecurso': ['pstotbjo', 'nombre recurso', 'recurso'],
    'Consecutivo': ['nrosecuencia', 'secuencia', 'operacion'],
    'Alistamiento': ['tiempominalistamientos', 'tiempo alistamiento'],
    'EquipoXRecurso': ['equiposxrecurso', 'equipos x recurso', 'equipo x recurso', 'cant equipos', 'cantequipos', '# equipos'],
    'OperarioXRecurso': ['totaloperarios', 'total operarios', 'mano de obra', 'operarios'],
    '%Enlace': ['porenlace', 'porcenlace', '% enlace', 'enlace'],
    '%Recorte': ['porcrecorte', '% recorte', 'recorte'],
    '%Sobrepeso': ['porcsobrepeso', '% sobrepeso', 'sobrepeso'],
    'AfectaEnlace': ['afectaenlace', 'afecta enlace'],
    'PESO GTA ROCEADA (gr)': ['pesogtaroceada', 'galleta roceada'],
    'Gtas/Sanduche': ['gtasanduche', 'galletas sanduche'],
    'Cant.crema sanduche': ['pesocremagta', 'cant crema sanduche'],
    'Galleta /paqueticos': ['gtaxpaquetico', 'galleta/paqueticos'],
    'PAQUETICOS UNIDAD': ['paqueticosxunidad', 'paqueticos unidad'],
    'UNIDADES CORRUGADO': ['unidadesxcorrugado', 'unidades corrugado'],
    'PESO N.IMPRESO UV(KG)': ['peso neto univenta especificaciones', 'peso n impreso']
}

def limpiar(texto):
    if pd.isna(texto): return ""
    t = str(texto)
    t = ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-zA-Z0-9]', '', t).lower()

def es_numero_real(val):
    if pd.isna(val) or str(val).strip() == "": return False
    try:
        float(str(val).replace(',', '.'))
        return True
    except ValueError:
        return False

def convertir_numero_normal(val):
    num = float(str(val).replace(',', '.'))
    return int(num) if num.is_integer() else num

def convertir_eficiencia(val):
    num = float(str(val).replace(',', '.'))
    if 0 < num <= 1.0: num = round(num * 100, 2)
    return int(num) if num.is_integer() else num

def resolver_operarios_y_equipos(recurso_str, dict_mo, dict_eq):
    rec = str(recurso_str).strip().upper()
    maquina = None
    if 'CAHO' in rec or 'CAM' in rec: maquina = 'CAMARA'
    elif 'EMBM' in rec or 'EMB' in rec: maquina = 'EMB'
    elif 'LAMI' in rec: maquina = 'LAMINA'
    elif 'EMPI' in rec or 'IND' in rec: maquina = 'IND'
    elif 'EMPM' in rec or 'MEP' in rec: maquina = 'MEP'
    elif 'ROCI' in rec: return 0.0, 1

    mo = dict_mo.get(maquina) if maquina else None
    eq = dict_eq.get(maquina) if maquina else None
    return mo, eq

def ajustar_secuencia_roceado(df_base):
    col_mat = next((c for c in df_base.columns if limpiar(c) == 'material'), None)
    col_rec = next((c for c in df_base.columns if limpiar(c) in MAPEO_BASE['NombreRecurso'] or limpiar(c) == 'nombrerecurso'), None)
    col_sec = next((c for c in df_base.columns if limpiar(c) in MAPEO_BASE['Consecutivo'] or limpiar(c) == 'consecutivo'), None)

    if not col_mat or not col_rec or not col_sec: return df_base

    for material, grupo in df_base.groupby(col_mat):
        recursos_material = grupo[col_rec].astype(str).str.upper().tolist()
        if any('ROCI' in r for r in recursos_material):
            for idx in grupo.index:
                maquina = str(df_base.at[idx, col_rec]).upper()
                if 'ROCI' in maquina: df_base.at[idx, col_sec] = 3
                elif 'EMPI' in maquina or 'IND' in maquina: df_base.at[idx, col_sec] = 4
                elif 'EMPM' in maquina or 'MEP' in maquina: df_base.at[idx, col_sec] = 5
                elif 'EMBM' in maquina or 'EMB' in maquina: df_base.at[idx, col_sec] = 6
    return df_base

def analizar_receta(xls):
    nombre_hoja = next((s for s in xls.sheet_names if 'receta' in s.lower()), xls.sheet_names[1])
    df_receta = pd.read_excel(xls, sheet_name=nombre_hoja, header=None)

    valores = {'eficiencia': None, 'ancho_molde': None, 'perimetro': None, 'velocidad': None,
               'largo_mm': None, 'ancho_mm': None, 'calibre': None, 'crudo': None, 'horneado': None}
    
    dict_mo = {'LAMINA': 3.5, 'CAMARA': 0.75, 'IND': 2.4, 'MEP': 2.3, 'EMB': 5.4}
    dict_eq = {'LAMINA': 1, 'CAMARA': 1, 'IND': 5, 'MEP': 2, 'EMB': 1}
    
    filas, columnas = df_receta.shape
    for r in range(filas):
        for c in range(columnas):
            celda = df_receta.iat[r, c]
            celda_limpia = limpiar(celda)
            
            if celda_limpia:
                if c + 1 < columnas:
                    valor = df_receta.iat[r, c + 1]
                    if es_numero_real(valor):
                        if 'eficiencia' in celda_limpia or 'efic' in celda_limpia: 
                            valores['eficiencia'] = convertir_eficiencia(valor)
                        else:
                            v_num = convertir_numero_normal(valor)
                            if 'anchomm' in celda_limpia or 'ancho(mm)' in celda_limpia: valores['ancho_mm'] = v_num
                            elif celda_limpia == 'ancho': valores['ancho_molde'] = v_num
                            elif 'perimetro' in celda_limpia: valores['perimetro'] = v_num
                            elif 'velocidad' in celda_limpia and 'camara' not in celda_limpia: valores['velocidad'] = v_num
                            elif 'largo' in celda_limpia: valores['largo_mm'] = v_num
                            elif 'calibre' in celda_limpia: valores['calibre'] = v_num
                            elif 'crudo' in celda_limpia or 'cruda' in celda_limpia: valores['crudo'] = v_num
                            elif 'horneado' in celda_limpia or 'horneada' in celda_limpia: valores['horneado'] = v_num

            val_upper = str(celda).strip().upper()
            if val_upper in ['LAMINA', 'CAMARA', 'IND', 'MEP', 'EMB']:
                if c + 1 < columnas and es_numero_real(df_receta.iat[r, c+1]): 
                    dict_mo[val_upper] = convertir_numero_normal(df_receta.iat[r, c+1])
                if c + 2 < columnas and es_numero_real(df_receta.iat[r, c+2]): 
                    dict_eq[val_upper] = convertir_numero_normal(df_receta.iat[r, c+2])

    return valores, dict_mo, dict_eq

# ==========================================
# INTERFAZ WEB (STREAMLIT)
# ==========================================
st.set_page_config(page_title="Actualizador SAP", layout="centered")

st.title("Actualizador Automatico SAP")
st.write("Sube el archivo Excel con la Base Original y la Receta. El sistema generara el Protocolo y las Especificaciones.")

archivo_subido = st.file_uploader("Selecciona el archivo Excel", type=["xlsx", "xls", "xlsm"])

if archivo_subido is not None:
    xls = pd.ExcelFile(archivo_subido)
    df_base = pd.read_excel(xls, sheet_name=0)
    
    valores_receta, dict_mo, dict_eq = analizar_receta(xls)
    
    st.success("Archivo leido correctamente")
    
    # Manejo de valor faltante sin ventanas de Windows
    if valores_receta['eficiencia'] is None:
        st.warning("No se detecto la Eficiencia de forma automatica en el archivo.")
        entrada_eficiencia = st.number_input("Ingresa el porcentaje de Eficiencia (Ej: 95):", min_value=0.0, max_value=100.0, value=90.0)
        valores_receta['eficiencia'] = convertir_eficiencia(entrada_eficiencia)

    if st.button("Procesar Archivo"):
        with st.spinner("Sincronizando equipos, mano de obra y secuencias..."):
            
            df_base_modificada = ajustar_secuencia_roceado(df_base.copy())
            columnas_limpias_base = {limpiar(col): col for col in df_base_modificada.columns}

            col_recurso_base = next((c for c in df_base_modificada.columns if limpiar(c) in MAPEO_BASE['NombreRecurso'] or limpiar(c) == 'nombrerecurso'), None)
            col_operarios_base = next((c for c in df_base_modificada.columns if limpiar(c) in MAPEO_BASE['OperarioXRecurso']), None)
            col_equipos_base = next((c for c in df_base_modificada.columns if limpiar(c) in MAPEO_BASE['EquipoXRecurso']), None)

            # --- 1. HOJA PROTOCOLO ---
            df_protocolo = pd.DataFrame(index=df_base_modificada.index)
            for col_sap in COLUMNAS_PROTOCOLO:
                col_busqueda = limpiar(col_sap)
                if col_busqueda in columnas_limpias_base:
                    df_protocolo[col_sap] = df_base_modificada[columnas_limpias_base[col_busqueda]].copy()
                elif col_sap in MAPEO_BASE:
                    encontrada = False
                    for alias in MAPEO_BASE[col_sap]:
                        if limpiar(alias) in columnas_limpias_base:
                            df_protocolo[col_sap] = df_base_modificada[columnas_limpias_base[limpiar(alias)]].copy()
                            encontrada = True
                            break
                    if not encontrada: df_protocolo[col_sap] = None
                else: df_protocolo[col_sap] = None

            if valores_receta['eficiencia'] is not None: df_protocolo['PorcEficiencia'] = valores_receta['eficiencia']
            if valores_receta['ancho_molde'] is not None: df_protocolo['Ancho'] = valores_receta['ancho_molde']
            if valores_receta['perimetro'] is not None: df_protocolo['Perimetro'] = valores_receta['perimetro']
            if valores_receta['velocidad'] is not None: df_protocolo['Velocidad'] = valores_receta['velocidad']

            for idx, fila in df_protocolo.iterrows():
                recurso = fila.get('NombreRecurso')
                mo, eq = resolver_operarios_y_equipos(recurso, dict_mo, dict_eq)
                
                if mo is not None: df_protocolo.at[idx, 'OperarioXRecurso'] = mo
                if eq is not None: df_protocolo.at[idx, 'EquipoXRecurso'] = eq
                
                if col_operarios_base and mo is not None: df_base_modificada.at[idx, col_operarios_base] = mo
                if col_equipos_base and eq is not None: df_base_modificada.at[idx, col_equipos_base] = eq

            # --- 2. HOJA ESPECIFICACIONES ---
            df_esp = pd.DataFrame(index=df_base_modificada.index)
            for col_sap in COLUMNAS_ESPECIFICACIONES:
                col_busqueda = limpiar(col_sap)
                if col_busqueda in columnas_limpias_base:
                    df_esp[col_sap] = df_base_modificada[columnas_limpias_base[col_busqueda]].copy()
                elif col_sap in MAPEO_BASE:
                    encontrada = False
                    for alias in MAPEO_BASE[col_sap]:
                        if limpiar(alias) in columnas_limpias_base:
                            df_esp[col_sap] = df_base_modificada[columnas_limpias_base[limpiar(alias)]].copy()
                            encontrada = True
                            break
                    if not encontrada: df_esp[col_sap] = None
                else: df_esp[col_sap] = None

            for col in df_base_modificada.columns:
                if limpiar(col) in ['material', 'codmaterial']: df_esp['MATERIAL'] = df_base_modificada[col].copy()

            if valores_receta['largo_mm'] is not None: df_esp['LARGO'] = valores_receta['largo_mm']
            if valores_receta['ancho_mm'] is not None: df_esp['ANCHO'] = valores_receta['ancho_mm']
            if valores_receta['calibre'] is not None: df_esp['CALIBRE'] = valores_receta['calibre']
            if valores_receta['crudo'] is not None: df_esp['Peso GTA cruda'] = valores_receta['crudo']
            if valores_receta['horneado'] is not None: df_esp['Peso GTA Horneada'] = valores_receta['horneado']

            if 'MATERIAL' in df_esp.columns: df_esp = df_esp.drop_duplicates(subset=['MATERIAL'])

            for idx, fila in df_protocolo.iterrows():
                for col_b, col_p in zip(['porceficiencia', 'ancho', 'perimetro', 'velocidad', 'totaloperarios', 'cantequipos'],
                                               ['PorcEficiencia', 'Ancho', 'Perimetro', 'Velocidad', 'OperarioXRecurso', 'EquipoXRecurso']):
                    if col_b in columnas_limpias_base:
                        df_base_modificada.at[idx, columnas_limpias_base[col_b]] = fila[col_p]

            # --- 3. GUARDADO EN MEMORIA Y FORMATO ---
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                ruta_tmp = tmp.name

            with pd.ExcelWriter(ruta_tmp, engine='openpyxl') as writer:
                df_protocolo.to_excel(writer, sheet_name='Protocolo', index=False)
                df_base_modificada.to_excel(writer, sheet_name='Base_Modificada', index=False)
                df_esp.to_excel(writer, sheet_name='Especificaciones', index=False)

            wb = openpyxl.load_workbook(ruta_tmp)
            
            wb['Base_Modificada'].sheet_properties.tabColor = "0000FF"
            amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

            for col_idx, nombre in enumerate(COLUMNAS_PROTOCOLO, 1):
                if nombre in ['Ancho', 'Perimetro', 'PorcEficiencia', 'EquipoXRecurso', 'Velocidad', 'OperarioXRecurso', 'AfectaEnlace', 'Consecutivo']:
                    wb['Protocolo'].cell(row=1, column=col_idx).fill = amarillo

            for col_idx, nombre in enumerate(COLUMNAS_ESPECIFICACIONES, 1):
                if nombre in ['Peso GTA cruda', 'Peso GTA Horneada', 'LARGO', 'ANCHO', 'CALIBRE']:
                    wb['Especificaciones'].cell(row=1, column=col_idx).fill = amarillo

            wb.save(ruta_tmp)

            with open(ruta_tmp, "rb") as file:
                datos_excel = file.read()

            # --- ANIMACIÓN Y BOTÓN DE DESCARGA ---
            st.balloons()
            st.success("Procesamiento completado. Tu archivo esta listo para descargar.")
            
            st.download_button(
                label="Descargar Actualizacion SAP",
                data=datos_excel,
                file_name="Carga_SAP_Actualizada.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
